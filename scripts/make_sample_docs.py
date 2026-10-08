"""Generate the fictional DOCX and PDF sample documents under docs/.

Run once: `uv run python scripts/make_sample_docs.py`. The PDF is written by hand
(plain Helvetica text objects) so the demo has no PDF-authoring dependency.
"""

from __future__ import annotations

import textwrap
from pathlib import Path

from docx import Document

DOCS = Path(__file__).resolve().parent.parent / "docs"

FAQ_SECTIONS: list[tuple[str, str]] = [
    (
        "Larkspur Hub Product FAQ - Overview",
        "Larkspur Hub is a fictional team knowledge platform used in this demo. This FAQ "
        "covers plans, API limits, security and support. For contractual terms refer to the "
        "Master Services Agreement.",
    ),
    (
        "Plans and Pricing",
        "Larkspur Hub offers three plans. The Starter plan costs USD 29 per month per workspace "
        "and includes up to 5 users. The Team plan costs USD 12 per user per month with a "
        "minimum of 5 users. The Enterprise plan is priced per contract. Annual billing is "
        "available on all plans and includes two months free compared with monthly billing. "
        "Upgrades take effect immediately and are prorated; downgrades take effect at the next "
        "renewal date.",
    ),
    (
        "API and Rate Limits",
        "Every plan includes REST API access. Rate limits are 120 requests per minute on "
        "Starter, 600 requests per minute on Team and 5,000 requests per minute on Enterprise. "
        "Requests above the limit receive HTTP 429 with a Retry-After header. Webhook deliveries "
        "are retried up to 5 times with exponential backoff over 24 hours. API keys can be "
        "rotated from Settings > Developer without downtime because two keys may be active at "
        "once.",
    ),
    (
        "Security and Data Retention",
        "Larkspur Hub is SOC 2 Type II certified. Customer data is stored in AWS regions "
        "eu-west-1 (Ireland) or us-east-1 (Virginia), selected when the workspace is created "
        "and not changeable afterwards. Single sign-on via SAML 2.0 is available on the Team "
        "and Enterprise plans. When a workspace is deleted its data is retained for 30 days "
        "and then permanently purged. Audit logs are kept for 400 days on Enterprise and 90 "
        "days on other plans.",
    ),
    (
        "Support",
        "Starter customers receive email support with a response target of 2 business days. "
        "Team customers receive email and chat support with a response target of 8 business "
        "hours. Enterprise customers receive 24/7 support with a 1 hour response target for "
        "critical incidents and a named customer success manager. Service status is published "
        "at status.larkspurhub.example.",
    ),
]

REFUND_PAGES: list[str] = [
    "Larkspur Hub Refund Policy\n"
    "Page 1 - Scope and money-back guarantee\n\n"
    "This policy applies to self-serve purchases of Larkspur Hub Starter and Team plans. It is "
    "a fictional document written for a demo project. Larkspur Hub offers a 30-day money-back "
    "guarantee on the first purchase of a monthly or annual plan. To qualify, the refund request "
    "must be received within 30 days of the first charge. Plan renewals, whether monthly or "
    "annual, are not eligible for the money-back guarantee. The guarantee may be used once per "
    "customer.",
    "Larkspur Hub Refund Policy\n"
    "Page 2 - Exclusions and partial refunds\n\n"
    "The following are not refundable: additional seats purchased mid-cycle, usage-based "
    "overage charges, and professional services. Enterprise contracts are governed by the "
    "Master Services Agreement and are excluded from this policy. Customers who cancel an "
    "annual plan after the 30-day window receive a prorated account credit for the remaining "
    "full months of the term. The credit is not paid out in cash and expires 12 months after "
    "it is issued.",
    "Larkspur Hub Refund Policy\n"
    "Page 3 - How to request a refund\n\n"
    "To request a refund, email billing@larkspurhub.example from the workspace owner account "
    "and include the workspace ID and the invoice number. Approved refunds are processed "
    "within 5 to 10 business days and are returned to the original payment method. Refunds "
    "cannot be redirected to a different card or bank account. Opening a chargeback instead of "
    "contacting billing will pause the workspace until the dispute is resolved.",
]


def make_docx() -> None:
    document = Document()
    for heading, body in FAQ_SECTIONS:
        document.add_heading(heading, level=1)
        document.add_paragraph(body)
    document.save(DOCS / "product_faq.docx")


def _pdf_escape(line: str) -> str:
    return line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _page_stream(text: str) -> bytes:
    lines: list[str] = []
    for paragraph in text.split("\n"):
        lines.extend(textwrap.wrap(paragraph, 90) or [""])
    body = "\n".join(f"({_pdf_escape(line)}) Tj T*" for line in lines)
    return f"BT /F1 11 Tf 14 TL 50 760 Td\n{body}\nET".encode("latin-1")


def make_pdf() -> None:
    objects: list[bytes] = []
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    page_ids = [4 + 2 * i for i in range(len(REFUND_PAGES))]
    kids = " ".join(f"{pid} 0 R" for pid in page_ids)
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i, text in enumerate(REFUND_PAGES):
        stream = _page_stream(text)
        content_id = page_ids[i] + 1
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>".encode()
        )
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    ).encode()
    (DOCS / "refund_policy.pdf").write_bytes(bytes(out))


if __name__ == "__main__":
    DOCS.mkdir(exist_ok=True)
    make_docx()
    make_pdf()
    print("wrote", DOCS / "product_faq.docx", "and", DOCS / "refund_policy.pdf")
